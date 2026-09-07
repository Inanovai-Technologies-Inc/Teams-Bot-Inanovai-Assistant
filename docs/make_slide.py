"""
Builds docs/workflow.pptx -- the Teams bot workflow as a PowerPoint slide.

Every box, arrow and label is a real PowerPoint shape, so you can click
and edit any of it. Nothing is a flat image.

Run:  python docs/make_slide.py
"""

import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x1B, 0x1F, 0x3B)
PURPLE = RGBColor(0x5B, 0x5F, 0xC7)
GREEN = RGBColor(0x10, 0x89, 0x3E)
AMBER = RGBColor(0xF2, 0xB7, 0x05)
BG = RGBColor(0xF5, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x60, 0x79)
BORDER = RGBColor(0xDD, 0xE0, 0xEE)
GREY = RGBColor(0x90, 0x96, 0xAE)
LIGHTGREY = RGBColor(0xB4, 0xB9, 0xCC)
MINT = RGBColor(0xE7, 0xF4, 0xEC)
LILAC = RGBColor(0xD6, 0xD8, 0xF2)
PALEGREEN = RGBColor(0xCF, 0xE8, 0xD8)
FONT = "Segoe UI"

# The design was drawn on a 1600x900 grid. The slide is 16:9, so one
# grid unit maps cleanly onto the slide at 120 units per inch.
UNITS_PER_INCH = 120


def u(v):
    """Grid units -> slide length."""
    return Inches(v / UNITS_PER_INCH)


def pt(v):
    """Grid units -> font points."""
    return Pt(v * 0.62)


prs = Presentation()
prs.slide_width = u(1600)
prs.slide_height = u(900)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def box(shape_type, x, y, w, h, fill=None, line=None, line_w=1.2,
        radius=None, dash=False):
    s = slide.shapes.add_shape(shape_type, u(x), u(y), u(w), u(h))
    if radius is not None:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # dashed
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.clear()
    return s


def label(x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
          spacing=None, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    for i, line_text in enumerate(text.split("\n")):
        para = p if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run()
        r.text = line_text
        r.font.size = pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color
        if spacing:
            # Letter spacing, in hundredths of a point. Plain attribute on
            # rPr, not namespaced, so it is set directly on the element.
            r.font._rPr.set("spc", str(int(spacing * 100)))
    return tb


# ---------------------------------------------------------------- canvas
box(MSO_SHAPE.RECTANGLE, 0, 0, 1600, 900, fill=BG)

# ---------------------------------------------------------------- header
box(MSO_SHAPE.RECTANGLE, 0, 0, 1600, 90, fill=NAVY)
box(MSO_SHAPE.RECTANGLE, 0, 87, 1600, 3, fill=PURPLE)
label(40, 22, 900, 34, "Microsoft Teams AI Bot", 28, WHITE, bold=True)
label(40, 60, 1000, 22,
      "System workflow — how a message travels from a Teams user to the AI model and back",
      14, RGBColor(0xAE, 0xB4, 0xD6))
badge = box(MSO_SHAPE.ROUNDED_RECTANGLE, 1394, 28, 166, 36, fill=PURPLE, radius=0.5)
label(1394, 39, 166, 20, "AI API  ·  Python", 13, WHITE, bold=True,
      align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- flow
label(40, 119, 300, 20, "MESSAGE FLOW", 12, PURPLE, bold=True, spacing=2)
box(MSO_SHAPE.RECTANGLE, 200, 123, 1360, 2, fill=BORDER)

STEPS = [
    ("1", "Teams User", "Types a message inside\nMicrosoft Teams.", PURPLE),
    ("2", "Microsoft Teams", "Delivers the message to\nthe registered bot.", PURPLE),
    ("3", "Azure Bot Service", "Authenticates the bot and\nroutes the message on.", PURPLE),
    ("4", "Dev Tunnel", "Public HTTPS bridge into\nthe local machine.", PURPLE),
    ("5", "app.py", "Web server. Receives the\nPOST on /api/messages.", PURPLE),
    ("6", "bot.py", "Teams layer. Strips the\nmention, blocks repeats.", PURPLE),
    ("7", "agent.py", "Agent layer. Adds memory\nand builds the prompt.", GREEN),
    ("8", "AI API", "Generates the reply and\nreturns it.", GREEN),
]

XS = [40, 432, 824, 1216]
CARD_W, CARD_H = 344, 118

for i, (num, title, body, accent) in enumerate(STEPS):
    row, col = divmod(i, 4)
    x = XS[col]
    y = 145 + row * 171

    box(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, CARD_W, CARD_H, fill=WHITE, radius=0.09)
    box(MSO_SHAPE.RECTANGLE, x, y, CARD_W, 4, fill=accent)
    box(MSO_SHAPE.OVAL, x + 17, y + 25, 30, 30, fill=accent)
    label(x + 17, y + 34, 30, 16, num, 13, WHITE, bold=True, align=PP_ALIGN.CENTER)
    label(x + 58, y + 33, CARD_W - 80, 22, title, 16, NAVY, bold=True)
    label(x + 26, y + 63, CARD_W - 44, 46, body, 12.5, MUTED)

    # arrow to the next card in the same row
    if col < 3:
        a = box(MSO_SHAPE.RIGHT_ARROW, x + CARD_W + 8, y + 50, 32, 18, fill=PURPLE)
        a.adjustments[0] = 0.5
        a.adjustments[1] = 0.55

# connector: card 4 down to card 5
box(MSO_SHAPE.RECTANGLE, 1386, 263, 3, 22, fill=PURPLE)
box(MSO_SHAPE.RECTANGLE, 212, 283, 1177, 3, fill=PURPLE)
box(MSO_SHAPE.RECTANGLE, 212, 283, 3, 14, fill=PURPLE)
down = box(MSO_SHAPE.DOWN_ARROW, 202, 292, 22, 22, fill=PURPLE)
down.adjustments[0] = 0.5
down.adjustments[1] = 0.5

# return path
box(MSO_SHAPE.RECTANGLE, 1386, 434, 3, 24, fill=GREEN)
ret = box(MSO_SHAPE.RECTANGLE, 212, 456, 1177, 3, fill=GREEN)
box(MSO_SHAPE.RECTANGLE, 212, 442, 3, 16, fill=GREEN)
up = box(MSO_SHAPE.UP_ARROW, 202, 424, 22, 22, fill=GREEN)
up.adjustments[0] = 0.5
up.adjustments[1] = 0.5
box(MSO_SHAPE.ROUNDED_RECTANGLE, 640, 441, 320, 34, fill=MINT, radius=0.5)
label(640, 450, 320, 20, "Reply returns along the same path", 13, GREEN,
      bold=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- safeguards
label(40, 513, 400, 20, "BUILT-IN SAFEGUARDS", 12, PURPLE, bold=True, spacing=2)
box(MSO_SHAPE.RECTANGLE, 270, 517, 1290, 2, fill=BORDER)

GUARDS = [
    ("Typing Indicator", "Refreshes every 4 seconds so\nthe chat never looks frozen."),
    ("Duplicate Guard", "Teams retries are ignored, so\nnobody is answered twice."),
    ("60 Second Timeout", "A stalled request fails politely\ninstead of hanging forever."),
    ("Echo Fallback", "Runs without an API key, so\nthe bot always starts."),
]

for i, (title, body) in enumerate(GUARDS):
    x = XS[i]
    box(MSO_SHAPE.ROUNDED_RECTANGLE, x, 540, CARD_W, 102, fill=WHITE, radius=0.1)
    box(MSO_SHAPE.RECTANGLE, x, 540, 5, 102, fill=AMBER)
    label(x + 28, 562, CARD_W - 50, 22, title, 14.5, NAVY, bold=True)
    label(x + 28, 592, CARD_W - 50, 44, body, 12, MUTED)

# ---------------------------------------------------------------- layers
label(40, 685, 300, 20, "LAYERED DESIGN", 12, PURPLE, bold=True, spacing=2)
box(MSO_SHAPE.RECTANGLE, 220, 689, 1340, 2, fill=BORDER)

box(MSO_SHAPE.ROUNDED_RECTANGLE, 40, 712, 880, 56, fill=PURPLE, radius=0.14)
label(66, 723, 830, 20, "Teams Layer — app.py and bot.py", 14.5, WHITE, bold=True)
label(66, 745, 830, 18, "Speaks Teams. Holds no intelligence. Never changes.", 12, LILAC)

box(MSO_SHAPE.ROUNDED_RECTANGLE, 40, 780, 880, 56, fill=GREEN, radius=0.14)
label(66, 791, 830, 20, "Agent Layer — agent.py", 14.5, WHITE, bold=True)
label(66, 813, 830, 18,
      "Holds the intelligence. Knows nothing about Teams. Fully swappable.", 12, PALEGREEN)

label(964, 723, 300, 18, "SWAPPABLE BRAIN", 12, MUTED, bold=True, spacing=1)

box(MSO_SHAPE.ROUNDED_RECTANGLE, 964, 752, 146, 48, fill=WHITE, line=BORDER, radius=0.16)
label(964, 762, 146, 18, "Echo Bot", 13.5, GREY, bold=True, align=PP_ALIGN.CENTER)
label(964, 781, 146, 14, "done", 10.5, LIGHTGREY, align=PP_ALIGN.CENTER)

a1 = box(MSO_SHAPE.RIGHT_ARROW, 1118, 767, 28, 18, fill=PURPLE)
a1.adjustments[0] = 0.5
a1.adjustments[1] = 0.55

box(MSO_SHAPE.ROUNDED_RECTANGLE, 1152, 746, 158, 60, fill=GREEN, radius=0.13)
label(1152, 760, 158, 18, "AI API", 13.5, WHITE, bold=True, align=PP_ALIGN.CENTER)
label(1152, 780, 158, 14, "live now", 10.5, PALEGREEN, align=PP_ALIGN.CENTER)

a2 = box(MSO_SHAPE.RIGHT_ARROW, 1318, 767, 28, 18, fill=PURPLE)
a2.adjustments[0] = 0.5
a2.adjustments[1] = 0.55

box(MSO_SHAPE.ROUNDED_RECTANGLE, 1352, 752, 148, 48, fill=WHITE, line=BORDER,
    radius=0.16, dash=True)
label(1352, 762, 148, 18, "NemoClaw", 13.5, GREY, bold=True, align=PP_ALIGN.CENTER)
label(1352, 781, 148, 14, "planned", 10.5, LIGHTGREY, align=PP_ALIGN.CENTER)

out = pathlib.Path(__file__).parent / "workflow.pptx"
prs.save(out)
print(f"Saved {out}")
print(f"Slide size: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches (16:9)")
print(f"Shapes on slide: {len(slide.shapes)}")
