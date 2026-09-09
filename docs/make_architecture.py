"""
Builds docs/architecture.pptx -- the multi-channel architecture as a slide.

Every box, arrow and label is a real PowerPoint shape, so you can click
and edit any of it. Nothing is a flat image.

Run:  python docs/make_architecture.py
"""

import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x1B, 0x1F, 0x3B)
PURPLE = RGBColor(0x5B, 0x5F, 0xC7)
GREEN = RGBColor(0x10, 0x89, 0x3E)
AMBER = RGBColor(0xF2, 0xB7, 0x05)
BG = RGBColor(0xF5, 0xF6, 0xFA)
PANEL = RGBColor(0xED, 0xEE, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x60, 0x79)
BORDER = RGBColor(0xDD, 0xE0, 0xEE)
OUTLINE = RGBColor(0xC6, 0xC9, 0xE0)
LILAC = RGBColor(0xAE, 0xB4, 0xD6)
PALEGREEN = RGBColor(0xCF, 0xE8, 0xD8)
FONT = "Segoe UI"

# Drawn on a 1600x900 grid, 120 units per inch, so the slide is 16:9.
UNITS_PER_INCH = 120


def u(v):
    return Inches(v / UNITS_PER_INCH)


def pt(v):
    return Pt(v * 0.62)


prs = Presentation()
prs.slide_width = u(1600)
prs.slide_height = u(900)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def box(shape_type, x, y, w, h, fill=None, line=None, line_w=1.2, radius=None):
    s = slide.shapes.add_shape(shape_type, u(x), u(y), u(w), u(h))
    if radius is not None:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.clear()
    return s


def label(x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
          spacing=None):
    tb = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    for i, line_text in enumerate(text.split("\n")):
        para = p if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run()
        r.text = line_text
        r.font.size = pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color
        if spacing:
            r.font._rPr.set("spc", str(int(spacing * 100)))
    return tb


def arrow_down(cx, y, h):
    body = box(MSO_SHAPE.RECTANGLE, cx - 1.5, y, 3, h - 14, fill=PURPLE)
    head = box(MSO_SHAPE.DOWN_ARROW, cx - 11, y + h - 20, 22, 20, fill=PURPLE)
    head.adjustments[0] = 0.5
    head.adjustments[1] = 0.5
    return body, head


COLS = [90, 590, 1090]
COL_W = 420
CENTRES = [300, 800, 1300]

# ---------------------------------------------------------------- canvas
box(MSO_SHAPE.RECTANGLE, 0, 0, 1600, 900, fill=BG)

# ---------------------------------------------------------------- header
box(MSO_SHAPE.RECTANGLE, 0, 0, 1600, 80, fill=NAVY)
box(MSO_SHAPE.RECTANGLE, 0, 77, 1600, 3, fill=PURPLE)
label(40, 18, 900, 32, "Inanovai Assistant — Architecture", 26, WHITE, bold=True)
label(40, 52, 900, 20, "Three chat apps, one bot, one brain", 14, LILAC)
box(MSO_SHAPE.ROUNDED_RECTANGLE, 1340, 22, 220, 34, fill=PURPLE, radius=0.5)
label(1340, 32, 220, 18, "Python  ·  Azure  ·  AI API", 13, WHITE, bold=True,
      align=PP_ALIGN.CENTER)

# ---------------------------------------------------- row 1: who is talking
label(40, 100, 300, 18, "WHO IS TALKING", 12, PURPLE, bold=True, spacing=2)
box(MSO_SHAPE.RECTANGLE, 230, 106, 1330, 2, fill=BORDER)

PEOPLE = [
    ("Microsoft Teams user", "live today"),
    ("Telegram user", "code ready, needs a token"),
    ("WhatsApp user", "code ready, needs Meta setup"),
]
for i, (title, note) in enumerate(PEOPLE):
    x = COLS[i]
    box(MSO_SHAPE.ROUNDED_RECTANGLE, x, 122, COL_W, 62, fill=WHITE, radius=0.16)
    box(MSO_SHAPE.RECTANGLE, x, 122, COL_W, 4, fill=PURPLE)
    label(x, 137, COL_W, 22, title, 17, NAVY, bold=True, align=PP_ALIGN.CENTER)
    label(x, 161, COL_W, 18, note, 12.5, MUTED, align=PP_ALIGN.CENTER)
    arrow_down(CENTRES[i], 184, 34)

# --------------------------------------------------------- row 2: platforms
label(40, 236, 300, 18, "THE POSTMAN", 12, PURPLE, bold=True, spacing=2)
box(MSO_SHAPE.RECTANGLE, 190, 242, 1370, 2, fill=BORDER)

PLATFORMS = [
    ("Azure Bot Service", "checks the bot is genuine,\nthen forwards the message"),
    ("Telegram Bot API", "posts every message to\nour webhook"),
    ("Meta Cloud API", "posts every message to\nour webhook"),
]
for i, (title, note) in enumerate(PLATFORMS):
    x = COLS[i]
    box(MSO_SHAPE.ROUNDED_RECTANGLE, x, 258, COL_W, 74, fill=WHITE, radius=0.14)
    box(MSO_SHAPE.RECTANGLE, x, 258, COL_W, 4, fill=PURPLE)
    label(x, 272, COL_W, 22, title, 16, NAVY, bold=True, align=PP_ALIGN.CENTER)
    label(x, 296, COL_W, 34, note, 12.5, MUTED, align=PP_ALIGN.CENTER)
    arrow_down(CENTRES[i], 332, 36)

# ------------------------------------------------------- row 3: our code
label(40, 384, 340, 18, "OUR CODE, ON AZURE", 12, PURPLE, bold=True, spacing=2)
box(MSO_SHAPE.RECTANGLE, 290, 390, 1270, 2, fill=BORDER)

box(MSO_SHAPE.ROUNDED_RECTANGLE, 66, 406, 1468, 200, fill=PANEL,
    line=OUTLINE, line_w=1.5, radius=0.08)
label(90, 420, 1000, 20,
      "Azure App Service — inanovai-teams-bot.azurewebsites.net", 13.5, PURPLE,
      bold=True)
label(90, 441, 1200, 18,
      "app.py listens on three addresses. Each one hands plain text to the same agent.",
      12, MUTED)

HANDLERS = [
    ("POST /api/messages", "bot.py",
     "strips the @mention,\nblocks repeat deliveries,\nshows the typing bubble"),
    ("POST /telegram/webhook", "channels/telegram.py",
     "checks the secret header,\nanswers only when mentioned\nin a group"),
    ("GET / POST /whatsapp/webhook", "channels/whatsapp.py",
     "answers Meta's verification,\nchecks the signature,\nmarks messages read"),
]
for i, (route, filename, note) in enumerate(HANDLERS):
    x = COLS[i]
    box(MSO_SHAPE.ROUNDED_RECTANGLE, x, 464, COL_W, 126, fill=WHITE, radius=0.09)
    box(MSO_SHAPE.RECTANGLE, x, 464, COL_W, 4, fill=PURPLE)
    label(x, 478, COL_W, 18, route, 13, PURPLE, bold=True, align=PP_ALIGN.CENTER)
    label(x, 500, COL_W, 22, filename, 16, NAVY, bold=True, align=PP_ALIGN.CENTER)
    label(x, 526, COL_W, 56, note, 12.5, MUTED, align=PP_ALIGN.CENTER)

# ------------------------------------------------ converge into the agent
for cx in CENTRES:
    box(MSO_SHAPE.RECTANGLE, cx - 1.5, 606, 3, 22, fill=GREEN)
box(MSO_SHAPE.RECTANGLE, 300, 625, 1000, 3, fill=GREEN)
box(MSO_SHAPE.RECTANGLE, 798.5, 625, 3, 18, fill=GREEN)
head = box(MSO_SHAPE.DOWN_ARROW, 789, 640, 22, 20, fill=GREEN)
head.adjustments[0] = 0.5
head.adjustments[1] = 0.5

# ------------------------------------------------------- row 4: the brain
box(MSO_SHAPE.ROUNDED_RECTANGLE, 440, 664, 720, 84, fill=GREEN, radius=0.12)
label(440, 678, 720, 24, "agent.py — the brain", 18, WHITE, bold=True,
      align=PP_ALIGN.CENTER)
label(440, 704, 720, 38,
      "adds the recent history for this conversation, builds the prompt\n"
      "knows nothing about Teams, Telegram or WhatsApp",
      12.5, PALEGREEN, align=PP_ALIGN.CENTER)

box(MSO_SHAPE.RECTANGLE, 798.5, 748, 3, 18, fill=GREEN)
head = box(MSO_SHAPE.DOWN_ARROW, 789, 762, 22, 20, fill=GREEN)
head.adjustments[0] = 0.5
head.adjustments[1] = 0.5

box(MSO_SHAPE.ROUNDED_RECTANGLE, 440, 786, 720, 70, fill=WHITE, radius=0.13)
box(MSO_SHAPE.RECTANGLE, 440, 786, 720, 4, fill=GREEN)
label(440, 802, 720, 24, "AI API — Llama 3.3 70B", 17, NAVY, bold=True,
      align=PP_ALIGN.CENTER)
label(440, 828, 720, 18, "writes the answer, which returns the way it came",
      12.5, MUTED, align=PP_ALIGN.CENTER)

# ------------------------------------------------------------- side panels
box(MSO_SHAPE.ROUNDED_RECTANGLE, 66, 664, 330, 120, fill=WHITE, radius=0.09)
box(MSO_SHAPE.RECTANGLE, 66, 664, 5, 120, fill=AMBER)
label(92, 680, 290, 20, "Chats stay separate", 14.5, NAVY, bold=True)
label(92, 706, 290, 36,
      "Every conversation gets its own\nmemory, labelled by platform:", 12, MUTED)
label(92, 744, 290, 32, "telegram:555\nwhatsapp:9199...", 12, PURPLE, bold=True)

box(MSO_SHAPE.ROUNDED_RECTANGLE, 1204, 664, 330, 120, fill=WHITE, radius=0.09)
box(MSO_SHAPE.RECTANGLE, 1204, 664, 5, 120, fill=AMBER)
label(1230, 680, 290, 20, "Same safeguards everywhere", 14.5, NAVY, bold=True)
label(1230, 708, 290, 70,
      "No duplicate answers\n60 second timeout\nLong replies split automatically\n"
      "Errors never crash a chat", 12, MUTED)

out = pathlib.Path(__file__).parent / "architecture.pptx"
prs.save(out)
print(f"Saved {out}")
print(f"Slide: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} in (16:9)")
print(f"Shapes: {len(slide.shapes)}")
